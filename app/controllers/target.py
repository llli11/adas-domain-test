import json
import logging
from typing import Any, Dict, Optional

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse

from app.models.ecu import ReleaseTargetInfo
from app.utils.get_ecu_full_info import (
    DESCRIPTION_TO_DID,
    DID_TO_DESCRIPTION,
    convert_description_based_to_did_based,
)

router = APIRouter()
logger = logging.getLogger(__name__)


def parse_excel_to_ecu_info(file_content: bytes) -> Dict[str, Any]:
    try:
        import io

        import openpyxl
    except ImportError:
        return {"error": "请安装 openpyxl 库: pip install openpyxl"}

    if not file_content or len(file_content) < 2:
        return {"error": "文件内容为空或过小"}

    try:
        if file_content[:2] == b"PK":
            workbook = openpyxl.load_workbook(io.BytesIO(file_content))
        elif file_content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            try:
                import xlrd

                workbook = xlrd.open_workbook(file_contents=file_content)
                sheet = workbook.sheet_by_index(0)
                ecu_data = {}
                for row_idx in range(1, sheet.nrows):
                    row = sheet.row_values(row_idx)
                    if not row or not row[0]:
                        continue
                    ecu_name = str(row[0])
                    ecu_item = {}
                    for i in range(1, len(row)):
                        if row[i]:
                            header = sheet.row_values(0)[i] if sheet.nrows > 0 and i < len(sheet.row_values(0)) else f"Field_{i}"
                            ecu_item[header] = str(row[i])
                    ecu_data[ecu_name] = ecu_item
                if not ecu_data:
                    return {"error": "Excel文件中没有有效数据"}
                return convert_description_based_to_did_based(ecu_data)
            except ImportError:
                return {"error": "请安装 xlrd 库解析旧版Excel文件: pip install xlrd"}
        else:
            try:
                text_content = file_content.decode("utf-8")
                csv_io = io.StringIO(text_content)
                import csv

                reader = list(csv.reader(csv_io))
                if not reader:
                    return {"error": "CSV文件为空"}

                headers = reader[0] if reader else []
                ecu_data = {}

                for row in reader[1:]:
                    if not row or not row[0]:
                        continue
                    ecu_name = row[0]
                    ecu_item = {}
                    for i in range(1, len(row)):
                        if row[i]:
                            if i < len(headers) and headers[i]:
                                header_name = headers[i]
                            else:
                                header_name = f"Field_{i}"
                            ecu_item[header_name] = row[i]
                    ecu_data[ecu_name] = ecu_item

                if not ecu_data:
                    return {"error": "文件中没有有效数据"}
                return convert_description_based_to_did_based(ecu_data)
            except UnicodeDecodeError as ue:
                return {
                    "error": f"文件格式错误: 不是有效的xlsx、xls或csv文件。请上传Excel文件(.xlsx, .xls)或CSV文件 (解码错误: {ue})"
                }

        sheet = workbook.active

        header_row = list(sheet.iter_rows(min_row=1, max_row=1, values_only=True))[0]
        headers = list(header_row) if header_row else []

        ecu_data = {}

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue

            ecu_name = row[0]
            ecu_item = {}

            for i in range(1, len(row)):
                if row[i]:
                    if i < len(headers) and headers[i]:
                        header_name = headers[i]
                    else:
                        header_name = f"Field_{i}"
                    ecu_item[header_name] = row[i]

            ecu_data[ecu_name] = ecu_item

        if not ecu_data:
            return {"error": "Excel文件中没有有效数据"}

        return convert_description_based_to_did_based(ecu_data)

    except Exception as e:
        return {"error": f"解析文件失败: {str(e)}"}


@router.get("/target/list", summary="获取基线列表")
async def get_target_list(search: Optional[str] = None) -> Dict[str, Any]:
    try:
        q = ReleaseTargetInfo.all().order_by("-updated_at")
        if search:
            q = ReleaseTargetInfo.filter(target_name__contains=search).order_by("-updated_at")
        results = await q.values("id", "target_name", "ecu_info", "created_at", "updated_at")
        return {"code": 200, "data": results, "msg": "OK"}
    except Exception as e:
        logger.error(f"Target list error: {e}")
        return {"code": 500, "data": [], "msg": f"数据库错误: {str(e)}"}


@router.get("/target/detail/{target_name}", summary="获取基线详情")
async def get_target_detail(target_name: str) -> Dict[str, Any]:
    record = await ReleaseTargetInfo.get_or_none(target_name=target_name)
    if not record:
        raise HTTPException(status_code=404, detail="未找到该基线版本的记录")
    return {
        "code": 200,
        "data": {
            "id": record.id,
            "target_name": record.target_name,
            "ecu_info": record.ecu_info,
            "created_at": record.created_at,
            "updated_at": record.updated_at,
        },
        "msg": "OK",
    }


@router.post("/target/update", summary="创建/更新基线")
async def update_target_info(target_name: str = Form(...), file: UploadFile = File(...)) -> Dict[str, Any]:
    file_content = await file.read()

    first_bytes = file_content[:8]

    if file_content[:2] == b"PK":
        ecu_data = parse_excel_to_ecu_info(file_content)
    elif len(file_content) > 100 and all(
        c in b"ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789+/=" for c in file_content[:100]
    ):
        import base64

        try:
            decoded = base64.b64decode(file_content)
            if decoded[:2] == b"PK":
                ecu_data = parse_excel_to_ecu_info(decoded)
            elif decoded[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
                import xlrd

                workbook = xlrd.open_workbook(file_contents=decoded)
                sheet = workbook.sheet_by_index(0)
                headers = sheet.row_values(0) if sheet.nrows > 0 else []
                ecu_data = {}
                for row_idx in range(1, sheet.nrows):
                    row = sheet.row_values(row_idx)
                    if not row or not row[0]:
                        continue
                    ecu_name = str(row[0])
                    ecu_item = {}
                    for i in range(1, len(row)):
                        if row[i]:
                            if i < len(headers) and headers[i]:
                                header_name = headers[i]
                            else:
                                header_name = f"Field_{i}"
                            ecu_item[header_name] = str(row[i])
                    ecu_data[ecu_name] = ecu_item
                if ecu_data:
                    ecu_data = convert_description_based_to_did_based(ecu_data)
                else:
                    ecu_data = None
            else:
                ecu_data = None
        except Exception:
            ecu_data = None
    elif first_bytes[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        try:
            import xlrd

            workbook = xlrd.open_workbook(file_contents=file_content)
            sheet = workbook.sheet_by_index(0)
            headers = sheet.row_values(0) if sheet.nrows > 0 else []
            ecu_data = {}
            for row_idx in range(1, sheet.nrows):
                row = sheet.row_values(row_idx)
                if not row or not row[0]:
                    continue
                ecu_name = str(row[0])
                ecu_item = {}
                for i in range(1, len(row)):
                    if row[i]:
                        if i < len(headers) and headers[i]:
                            header_name = headers[i]
                        else:
                            header_name = f"Field_{i}"
                        ecu_item[header_name] = str(row[i])
                ecu_data[ecu_name] = ecu_item
            if not ecu_data:
                raise HTTPException(status_code=400, detail="Excel文件中没有有效数据")
            ecu_data = convert_description_based_to_did_based(ecu_data)
        except ImportError:
            raise HTTPException(status_code=400, detail="请安装 xlrd 库解析旧版Excel文件")
    else:
        ecu_data = None
        encodings = ["utf-8", "gbk", "gb2312", "latin1", "cp1252"]
        for enc in encodings:
            try:
                decoded = file_content.decode(enc)
                import csv
                import io

                csv_io = io.StringIO(decoded)
                reader = list(csv.reader(csv_io))
                if reader:
                    temp_data = {}
                    for row in reader[1:]:
                        if not row or not row[0]:
                            continue
                        ecu_name = str(row[0])
                        ecu_item = {}
                        for i in range(1, len(row)):
                            if row[i]:
                                ecu_item[f"Field_{i}"] = str(row[i])
                        temp_data[ecu_name] = ecu_item
                    if temp_data:
                        ecu_data = convert_description_based_to_did_based(temp_data)
                        break
            except Exception:
                continue

        if ecu_data is None:
            pk_index = file_content.find(b"PK")
            if pk_index > 0 and pk_index < 100:
                file_content = file_content[pk_index:]
                ecu_data = parse_excel_to_ecu_info(file_content)
            else:
                raise HTTPException(
                    status_code=400,
                    detail=f"文件格式无法识别。文件头字节: {file_content[:20].hex()}。请确保上传的是有效的 Excel 文件(.xlsx, .xls)",
                )

    if ecu_data is None:
        raise HTTPException(
            status_code=400,
            detail=f"文件格式无法识别。文件头字节: {file_content[:20].hex()}。请确保上传的是有效的 Excel 文件(.xlsx, .xls)",
        )

    if isinstance(ecu_data, dict) and "error" in ecu_data:
        raise HTTPException(status_code=400, detail=ecu_data["error"])

    existing = await ReleaseTargetInfo.get_or_none(target_name=target_name)

    if existing:
        existing.ecu_info = ecu_data
        await existing.save()
    else:
        await ReleaseTargetInfo.create(target_name=target_name, ecu_info=ecu_data)

    return {"code": 200, "data": {"message": "更新成功", "target_name": target_name}, "msg": "OK"}


@router.delete("/target/delete/{target_name}", summary="删除基线")
async def delete_target_info(target_name: str) -> Dict[str, Any]:
    await ReleaseTargetInfo.filter(target_name=target_name).delete()
    return {"code": 200, "data": {"message": "删除成功"}, "msg": "OK"}


@router.get("/target/template", summary="获取模板")
async def get_template():
    try:
        import io

        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "基线版本ECU"

        headers = [
            "ECU名",
            "VOYAH SoftwareVersion",
            "VOYAH HardwareVersion",
            "Supplier SoftwareVersion",
            "Supplier HardwareVersion",
        ]
        for col, header in enumerate(headers, 1):
            sheet.cell(1, col, header)

        example_data = [
            ["ACU", "H77A3607802AE", "H77A3607801AA", "SW_DFM_H77_BL1.00", "H77A16FL"],
            ["BMS", "H56E3619020AB", "H56E3619019AA", "A0BMUQ1B5D14AV06", "H56E3619"],
        ]
        for row_idx, row_data in enumerate(example_data, 2):
            for col_idx, value in enumerate(row_data, 1):
                sheet.cell(row_idx, col_idx, value)

        for col in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 25

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)

        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=template.xlsx"},
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"生成模板失败: {str(e)}")
