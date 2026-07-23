import os
from datetime import datetime, date
from io import BytesIO

from fastapi import APIRouter, Query, UploadFile, File, Body, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook

from app.controllers.tool import (
    tool_controller,
    tool_borrow_controller,
    tool_inventory_controller,
    tool_inventory_detail_controller,
    tool_requirement_controller,
)
from app.core.dependency import DependAuth
from app.schemas import Success, Fail
from app.settings import settings
from app.schemas.tool import (
    ToolCreate,
    ToolUpdate,
    ToolBorrowCreate,
    ToolBorrowUpdate,
    ToolInventoryCreate,
    ToolInventoryUpdate,
    ToolInventoryDetailCreate,
    ToolRequirementCreate,
    ToolRequirementUpdate,
)

router = APIRouter(tags=["工具管理"])


# ==================== 工具台账管理 ====================
@router.get("/list", summary="查看工具列表")
async def list_tools(
    tool_code: str = Query(None, description="工具编码"),
    tool_name: str = Query(None, description="工具名称"),
    tool_type: str = Query(None, description="工具类型"),
    status: str = Query(None, description="状态"),
    current_user: str = Query(None, description="当前使用者"),
):
    tools = await tool_controller.search_tools(
        tool_code=tool_code, tool_name=tool_name, tool_type=tool_type, status=status, current_user=current_user
    )
    data = [await t.to_dict() for t in tools]
    return Success(data=data)


@router.get("/get", summary="查看工具详情")
async def get_tool(id: int = Query(..., description="工具ID")):
    tool_obj = await tool_controller.get(id=id)
    data = await tool_obj.to_dict()
    return Success(data=data)


@router.post("/create", summary="创建工具")
async def create_tool(tool_in: ToolCreate):
    exist = await tool_controller.model.filter(tool_code=tool_in.tool_code).first()
    if exist:
        return Fail(msg="工具编码已存在")
    await tool_controller.create(obj_in=tool_in)
    return Success(msg="创建成功")


@router.post("/update", summary="更新工具")
async def update_tool(id: int = Query(..., description="工具ID"), tool_in: ToolUpdate = None):
    await tool_controller.update(id=id, obj_in=tool_in)
    return Success(msg="更新成功")


@router.delete("/delete", summary="删除工具")
async def delete_tool(id: int = Query(..., description="工具ID")):
    await tool_controller.remove(id=id)
    return Success(msg="删除成功")


@router.delete("/batch-delete", summary="批量删除工具")
async def batch_delete_tools():
    count = await tool_controller.model.all().delete()
    return Success(msg=f"已删除 {count} 条数据")


# ==================== 工具借用管理 ====================
@router.get("/borrow/list", summary="查看借用列表")
async def list_borrows(
    tool_code: str = Query(None, description="工具编码"),
    tool_name: str = Query(None, description="工具名称"),
    borrower_name: str = Query(None, description="借用人姓名"),
    status: str = Query(None, description="借用状态"),
    approve_status: str = Query(None, description="审批状态"),
):
    borrows = await tool_borrow_controller.search_borrows(
        tool_code=tool_code,
        tool_name=tool_name,
        borrower_name=borrower_name,
        status=status,
        approve_status=approve_status,
    )
    data = [await b.to_dict() for b in borrows]
    return Success(data=data)


@router.get("/borrow/get", summary="查看借用详情")
async def get_borrow(id: int = Query(..., description="借用记录ID")):
    borrow_obj = await tool_borrow_controller.get(id=id)
    data = await borrow_obj.to_dict()
    return Success(data=data)


@router.post("/borrow/create", summary="创建借用申请")
async def create_borrow(borrow_in: ToolBorrowCreate):
    tool = await tool_controller.get(id=borrow_in.tool_id)
    if tool.status != "可用":
        return Fail(msg="该工具当前不可借用")
    await tool_borrow_controller.create(obj_in=borrow_in)
    return Success(msg="申请成功")


@router.post("/borrow/approve", summary="审批借用申请")
async def approve_borrow(
    id: int = Body(..., description="借用记录ID"),
    approve_status: str = Body(..., description="审批状态"),
    approver_id: int = Body(None, description="审批人ID"),
    approver_name: str = Body(None, description="审批人姓名"),
    current_user = DependAuth,
):
    borrow = await tool_borrow_controller.get(id=id)
    if borrow.approve_status != "待审批":
        return Fail(msg="该申请已审批")
    
    # 自动从当前登录用户填充审批人信息
    approver_id = approver_id or current_user.id
    approver_name = approver_name or current_user.username
    
    update_data = ToolBorrowUpdate(
        approve_status=approve_status,
        approver_id=approver_id,
        approver_name=approver_name,
        approve_time=datetime.now(),
    )
    await tool_borrow_controller.update(id=id, obj_in=update_data)
    
    if approve_status == "已通过":
        tool = await tool_controller.get(id=borrow.tool_id)
        await tool_controller.update(id=tool.id, obj_in=ToolUpdate(status="借出"))
    
    return Success(msg="审批成功")


@router.post("/borrow/return", summary="归还工具")
async def return_borrow(id: int = Body(..., description="借用记录ID")):
    borrow = await tool_borrow_controller.get(id=id)
    if borrow.status != "借用中":
        return Fail(msg="该工具已归还")
    
    update_data = ToolBorrowUpdate(
        status="已归还",
        actual_return_date=datetime.now(),
    )
    await tool_borrow_controller.update(id=id, obj_in=update_data)
    
    tool = await tool_controller.get(id=borrow.tool_id)
    await tool_controller.update(id=tool.id, obj_in=ToolUpdate(status="可用"))
    
    return Success(msg="归还成功")


# ==================== 工具盘点管理 ====================
@router.get("/inventory/list", summary="查看盘点任务列表")
async def list_inventories(
    task_code: str = Query(None, description="任务编码"),
    task_name: str = Query(None, description="任务名称"),
    status: str = Query(None, description="状态"),
):
    inventories = await tool_inventory_controller.search_inventories(
        task_code=task_code, task_name=task_name, status=status
    )
    data = [await i.to_dict() for i in inventories]
    return Success(data=data)


@router.get("/inventory/get", summary="查看盘点任务详情")
async def get_inventory(id: int = Query(..., description="盘点任务ID")):
    inventory_obj = await tool_inventory_controller.get(id=id)
    data = await inventory_obj.to_dict()
    details = await tool_inventory_detail_controller.get_details_by_inventory(inventory_id=id)
    data["details"] = [await d.to_dict() for d in details]
    return Success(data=data)


@router.post("/inventory/create", summary="创建盘点任务")
async def create_inventory(inventory_in: ToolInventoryCreate):
    exist = await tool_inventory_controller.model.filter(task_code=inventory_in.task_code).first()
    if exist:
        return Fail(msg="任务编码已存在")
    
    inventory = await tool_inventory_controller.create(obj_in=inventory_in)
    
    tools = await tool_controller.model.all()
    details = []
    for tool in tools:
        details.append(
            ToolInventoryDetailCreate(
                inventory_id=inventory.id,
                tool_id=tool.id,
                tool_code=tool.tool_code,
                tool_name=tool.tool_name,
                book_quantity=tool.quantity,
                actual_quantity=0,
                diff_quantity=-tool.quantity,
                status="待盘点",
            )
        )
    
    for detail in details:
        await tool_inventory_detail_controller.create(obj_in=detail)
    
    update_data = ToolInventoryUpdate(total_count=len(tools))
    await tool_inventory_controller.update(id=inventory.id, obj_in=update_data)
    
    return Success(msg=f"盘点任务已创建，共{len(tools)}件工具")


@router.post("/inventory/complete", summary="完成盘点")
async def complete_inventory(
    id: int = Body(..., description="盘点任务ID"),
    actual_count: int = Body(..., description="实盘数量"),
    diff_explanation: str = Body("", description="差异说明"),
):
    inventory = await tool_inventory_controller.get(id=id)
    if inventory.status == "已完成":
        return Fail(msg="该任务已完成")
    
    details = await tool_inventory_detail_controller.get_details_by_inventory(inventory_id=id)
    diff_count = inventory.total_count - actual_count
    
    update_data = ToolInventoryUpdate(
        status="已完成",
        actual_count=actual_count,
        diff_count=diff_count,
        diff_explanation=diff_explanation,
        complete_time=datetime.now(),
    )
    await tool_inventory_controller.update(id=id, obj_in=update_data)
    
    return Success(msg="盘点完成")


@router.put("/inventory/update", summary="更新盘点任务")
async def update_inventory(
    id: int = Body(..., description="盘点任务ID"),
    actual_count: int = Body(..., description="实盘数量"),
    diff_explanation: str = Body("", description="差异说明"),
):
    inventory = await tool_inventory_controller.get(id=id)
    diff_count = inventory.total_count - actual_count
    
    update_data = ToolInventoryUpdate(
        actual_count=actual_count,
        diff_count=diff_count,
        diff_explanation=diff_explanation,
    )
    await tool_inventory_controller.update(id=id, obj_in=update_data)
    
    return Success(msg="更新成功")


@router.delete("/inventory/delete", summary="删除盘点任务")
async def delete_inventory(id: int = Query(..., description="盘点任务ID")):
    inventory = await tool_inventory_controller.get(id=id)
    if inventory.status == "已完成":
        return Fail(msg="已完成的盘点任务不能删除")
    
    await tool_inventory_detail_controller.model.filter(inventory_id=id).delete()
    await tool_inventory_controller.model.filter(id=id).delete()
    
    return Success(msg="删除成功")


# ==================== 工具需求管理 ====================
@router.get("/requirement/list", summary="查看需求列表")
async def list_requirements(
    tool_name: str = Query(None, description="工具名称"),
    tool_type: str = Query(None, description="工具类型"),
    requester_name: str = Query(None, description="需求人姓名"),
    status: str = Query(None, description="状态"),
):
    requirements = await tool_requirement_controller.search_requirements(
        tool_name=tool_name,
        tool_type=tool_type,
        requester_name=requester_name,
        status=status,
    )
    data = [await r.to_dict() for r in requirements]
    return Success(data=data)


@router.get("/requirement/get", summary="查看需求详情")
async def get_requirement(id: int = Query(..., description="需求ID")):
    requirement_obj = await tool_requirement_controller.get(id=id)
    data = await requirement_obj.to_dict()
    return Success(data=data)


@router.post("/requirement/create", summary="创建需求申请")
async def create_requirement(requirement_in: ToolRequirementCreate):
    await tool_requirement_controller.create(obj_in=requirement_in)
    return Success(msg="申请成功")


@router.post("/requirement/handle", summary="处理需求申请")
async def handle_requirement(
    id: int = Body(..., description="需求ID"),
    status: str = Body(..., description="处理状态"),
    handler_id: int = Body(..., description="处理人ID"),
    handler_name: str = Body(..., description="处理人姓名"),
):
    requirement = await tool_requirement_controller.get(id=id)
    if requirement.status != "待处理":
        return Fail(msg="该需求已处理")
    
    update_data = ToolRequirementUpdate(
        status=status,
        handler_id=handler_id,
        handler_name=handler_name,
        handle_time=datetime.now(),
    )
    await tool_requirement_controller.update(id=id, obj_in=update_data)
    
    return Success(msg="处理成功")


@router.put("/requirement/update", summary="更新需求")
async def update_requirement(
    id: int = Body(..., description="需求ID"),
    tool_name: str = Body(..., description="工具名称"),
    spec_model: str = Body("", description="规格型号"),
    quantity: int = Body(..., description="需求数量"),
    requester_name: str = Body(..., description="需求人姓名"),
    requirement_date: date = Body(..., description="需求日期"),
    reason: str = Body("", description="需求原因"),
):
    requirement = await tool_requirement_controller.get(id=id)
    if requirement.status != "待处理":
        return Fail(msg="已处理的需求不能修改")
    
    update_data = ToolRequirementUpdate(
        tool_name=tool_name,
        spec_model=spec_model,
        quantity=quantity,
        requester_name=requester_name,
        requirement_date=requirement_date,
        reason=reason,
    )
    await tool_requirement_controller.update(id=id, obj_in=update_data)
    
    return Success(msg="更新成功")


@router.delete("/requirement/delete", summary="删除需求")
async def delete_requirement(id: int = Query(..., description="需求ID")):
    requirement = await tool_requirement_controller.get(id=id)
    if requirement.status != "待处理":
        return Fail(msg="已处理的需求不能删除")
    
    await tool_requirement_controller.model.filter(id=id).delete()
    
    return Success(msg="删除成功")


# ==================== 图片上传 ====================
# 图片访问路由——独立 router，无需认证（浏览器 <img> 标签不带 token）
image_router = APIRouter(tags=["工具管理"])


@image_router.get("/{filename}", summary="获取设备图片")
async def get_image(filename: str):
    from fastapi.responses import FileResponse
    from fastapi import HTTPException
    from app.settings import settings
    upload_dir = os.path.join(settings.BASE_DIR, "app", "static", "uploads")
    file_path = os.path.join(upload_dir, filename)
    if not os.path.isfile(file_path):
        raise HTTPException(status_code=404, detail="图片不存在")
    return FileResponse(file_path)


@router.post("/upload", summary="上传设备图片")
async def upload_image(file: UploadFile = File(None)):
    if not file or not file.filename:
        return Fail(msg="请选择图片文件")
    if not file.content_type or not file.content_type.startswith("image/"):
        return Fail(msg="请上传图片文件")
    
    base_dir = settings.BASE_DIR
    upload_dir = os.path.join(base_dir, "app", "static", "uploads")
    os.makedirs(upload_dir, exist_ok=True)
    
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    ext = file.filename.split(".")[-1] if "." in file.filename else "png"
    filename = f"{timestamp}.{ext}"
    filepath = os.path.join(upload_dir, filename)
    
    contents = await file.read()
    with open(filepath, "wb") as f:
        f.write(contents)
    
    url = f"/api/v1/tool/image/{filename}"
    return Success(data={"url": url}, msg="上传成功")

# ==================== 数据导入导出 ====================
@router.get("/export", summary="导出工具数据(Excel)")
async def export_tools():
    tools = await tool_controller.model.all()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "工具台账"
    
    headers = ["设备编号", "设备名称", "设备类别", "设备数量", "设备图片", "当前使用者", "状态", "备注"]
    ws.append(headers)
    
    for tool in tools:
        row = [
            tool.tool_code,
            tool.tool_name,
            tool.tool_type,
            tool.quantity,
            tool.image_url or "",
            tool.current_user or "",
            tool.status or "可用",
            tool.remark or "",
        ]
        ws.append(row)
    
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    response = StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response.headers["Content-Disposition"] = "attachment; filename=tools.xlsx"
    return response


@router.post("/import", summary="导入工具数据(Excel)")
async def import_tools(file: UploadFile = File(...)):
    try:
        content = await file.read()
        wb = load_workbook(BytesIO(content))
        ws = wb.active
        
        # 读取表头行，建立列名到索引的映射
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header_row:
            return Fail(msg="Excel 文件格式错误：无法读取表头行")
        
        # 定义标准列名映射
        col_map = {}
        header_names = {
            "设备编号": "tool_code",
            "设备名称": "tool_name",
            "设备类别": "tool_type",
            "设备数量": "quantity",
            "设备图片": "image_url",
            "当前使用者": "current_user",
            "状态": "status",
            "备注": "remark",
        }
        for idx, cell in enumerate(header_row):
            if cell and str(cell).strip() in header_names:
                col_map[header_names[str(cell).strip()]] = idx
        
        # 校验必要的列是否存在
        required_cols = ["tool_code", "tool_name"]
        for col in required_cols:
            if col not in col_map:
                return Fail(msg=f"Excel 文件缺少必要列：{list(header_names.keys())[list(header_names.values()).index(col)]}")
        
        def get_cell_value(row, field, default=None):
            """根据字段名从行数据中取值"""
            idx = col_map.get(field)
            if idx is None:
                return default
            if idx >= len(row):
                return default
            val = row[idx]
            if val is None:
                return default
            cleaned = str(val).strip()
            if cleaned == "None":
                return default
            return cleaned
        
        success_count = 0
        fail_count = 0
        errors = []
        
        for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            tool_code = get_cell_value(row, "tool_code")
            tool_name = get_cell_value(row, "tool_name")
            
            if not tool_code or not tool_name:
                continue
            
            try:
                exist = await tool_controller.model.filter(tool_code=tool_code).first()
                if exist:
                    fail_count += 1
                    errors.append(f"第{row_num}行：设备编号 {tool_code} 已存在")
                    continue
                
                tool_type = get_cell_value(row, "tool_type", "")
                quantity_raw = get_cell_value(row, "quantity")
                try:
                    quantity = int(float(quantity_raw)) if quantity_raw else 1
                except (ValueError, TypeError):
                    quantity = 1
                
                image_url = get_cell_value(row, "image_url")
                current_user = get_cell_value(row, "current_user")
                
                status = get_cell_value(row, "status", "可用")
                remark = get_cell_value(row, "remark", "")
                
                # 根据状态判断是否在库
                is_in_stock = status not in ("借出", "报废")
                
                await tool_controller.create(obj_in=ToolCreate(
                    tool_code=tool_code,
                    tool_name=tool_name,
                    tool_type=tool_type,
                    quantity=quantity,
                    image_url=image_url,
                    current_user=current_user,
                    status=status,
                    is_in_stock=is_in_stock,
                    remark=remark,
                ))
                success_count += 1
            except Exception as e:
                fail_count += 1
                errors.append(f"第{row_num}行：{str(e)}")
        
        msg = f"导入完成，成功 {success_count} 条，失败 {fail_count} 条"
        if errors:
            msg += "\n" + "\n".join(errors[:10])
        return Success(msg=msg)
    except Exception as e:
        return Fail(msg=f"导入失败：{str(e)}")
