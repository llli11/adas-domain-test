"""驾驶员月度结算"""
import calendar
from datetime import date
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.expense import DailyRecord, SupplierRate
from app.schemas.base import Success

router = APIRouter(tags=["费用管理"])

SERVICE_LEVEL = "不涉及"
SERVICE_DESC = "智驾域测试业务支持"


def _get_attribute(rec):
    return "加班" if rec.is_overtime == "是" else "正常"


def _get_travel_status(rec):
    """出差状态：出差 / 未出差"""
    ts = rec.travel_status or ""
    return "出差" if ts == "出差" else "未出差"


async def _build_grouped_data(records, days_in_month):
    """按 (person_name, attribute, travel_status, supplier) 分组，正常用 normal_hours，加班用 overtime_hours。
    数据源为 DailyRecord，supplier 直接取记录自带值。"""
    # 第一步：聚合原始数据
    grouped = {}
    all_suppliers = {}  # person_name → supplier

    # 预建 test_order_id → supplier（从非空 supplier 的记录），供 supplier 空时按同试验单兜底
    to_supplier_map = {}
    for rec in records:
        s = (rec.supplier or "").strip()
        _toid = rec.test_order_id if rec.test_order_id is not None else (rec.test_order.id if rec.test_order else None)
        if s and _toid:
            to_supplier_map.setdefault(_toid, s)
        if "樊友康" in (rec.person_name or ""):
            print(f"[DriverSettlement] 樊友康: test_order_id={rec.test_order_id!r}, test_order={rec.test_order!r}, supplier={rec.supplier!r}, to_map={to_supplier_map}")

    for rec in records:
        supplier = (rec.supplier or "").strip()
        _toid = rec.test_order_id if rec.test_order_id is not None else (rec.test_order.id if rec.test_order else None)
        if not supplier and _toid:
            supplier = to_supplier_map.get(_toid, "")
        if not supplier:
            continue  # 没有供应商，跳过

        all_suppliers[rec.person_name] = supplier
        tstatus = _get_travel_status(rec)
        day = rec.record_date.day

        # 一条记录拆成两行：正常取 normal_hours，加班取 overtime_hours
        work_dur = float(rec.normal_hours or 0)
        overtime_dur = float(rec.overtime_hours or 0)

        # 正常行（总是需要，work_duration > 0 时填充）
        key_normal = (rec.person_name, supplier, "正常", tstatus)
        if key_normal not in grouped:
            grouped[key_normal] = {}
        if work_dur > 0:
            grouped[key_normal][day] = grouped[key_normal].get(day, 0) + work_dur

        # 加班行（overtime_hours > 0 时填充）
        key_overtime = (rec.person_name, supplier, "加班", tstatus)
        if key_overtime not in grouped:
            grouped[key_overtime] = {}
        if overtime_dur > 0:
            grouped[key_overtime][day] = grouped[key_overtime].get(day, 0) + overtime_dur

    # 第二步：确保每人都有4种组合（正常/未出差、加班/未出差、正常/出差、加班/出差）
    # 预加载 SupplierRate 按 name（合同号/编号为供应商级，各时段一致，取第一行）
    sr_map = {}
    for sr in await SupplierRate.all():
        if sr.name not in sr_map:
            sr_map[sr.name] = sr

    records_list = []
    person_summary = {}

    for person_name, supplier in all_suppliers.items():
        sr = sr_map.get(supplier)
        if not sr:
            continue  # 供应商单价表无此供应商，跳过
        contract_no = sr.contract_no or ""

        # 遍历4种组合
        for attr in ["正常", "加班"]:
            for tstatus in ["未出差", "出差"]:
                code = (sr.code_local if tstatus == "未出差" else sr.code_trip) or ""
                service_name = f"{code} {SERVICE_DESC}" if code else ""

                key = (person_name, supplier, attr, tstatus)
                day_map = grouped.get(key, {})

                row_days = 0
                row_hours = 0.0
                days_dict = {}
                for d in range(1, days_in_month + 1):
                    hours = day_map.get(d, 0)
                    if hours > 0:
                        days_dict[str(d)] = round(hours / 8, 3)
                        row_days += 1
                        row_hours += hours

                row = {
                    "person_name": person_name,
                    "contract_no": contract_no,
                    "service_category": service_name,
                    "service_level": SERVICE_LEVEL,
                    "attribute": attr,
                    "travel_status": tstatus,
                    "days": days_dict,
                    "service_days": row_days,
                    "service_hours": round(row_hours, 3),
                    "remark": "",
                }
                records_list.append(row)

                # 汇总（按人+出差状态分组）
                key_ts = (person_name, tstatus)
                if key_ts not in person_summary:
                    person_summary[key_ts] = {
                        "normal_hours": 0.0, "overtime_hours": 0.0, "contract_no": contract_no,
                        "person_name": person_name, "travel_status": tstatus,
                    }
                if attr == "正常":
                    person_summary[key_ts]["normal_hours"] += row_hours
                else:
                    person_summary[key_ts]["overtime_hours"] += row_hours

    # 排序：按人名 → 未出差优先(RWYF0549072) → 正常优先
    records_list.sort(key=lambda r: (
        r["person_name"],
        0 if r["travel_status"] == "未出差" else 1,
        0 if r["attribute"] == "正常" else 1,
    ))

    return records_list, person_summary


# ==================== 导出 Excel ====================

@router.get("/export", summary="导出驾驶员月度结算Excel")
async def export_driver_settlement(
    year_month: str = Query(..., description="结算月份(YYYY-MM)"),
    test_order_id: int = Query(None, description="试验单号ID"),
):
    year, month = int(year_month[:4]), int(year_month[5:7])
    days_in_month = calendar.monthrange(year, month)[1]

    month_start = date(year, month, 1)
    month_end = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)

    from tortoise.expressions import Q
    q = (
        Q(record_date__gte=str(month_start), record_date__lt=str(month_end))
        & Q(person_type="驾驶员")
        & ~Q(approval_status="驳回")
    )
    if test_order_id and isinstance(test_order_id, int) and test_order_id > 0:
        q &= Q(test_order_id=test_order_id)

    records = await DailyRecord.filter(q).select_related("test_order").order_by("person_name", "record_date")
    records_list, person_summary = await _build_grouped_data(records, days_in_month)

    # 样式
    title_font = Font(name="Arial", size=12, bold=True)
    header_font = Font(name="Arial", size=9, bold=True)
    cell_font = Font(name="Arial", size=9)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    wb = Workbook()
    ws = wb.active
    ws.title = f"{year}年{month}月驾驶员月度结算"

    # 列：NO. | 合同号 | 服务人员姓名 | 服务类别 | 服务等级 | 属性 | 日期1-31 | 备注 | 合计正常服务天数 | 合计额外服务天数 | 合计总服务天数
    BASE_COLS = 6
    DATE_START = BASE_COLS + 1
    REMARK_COL = DATE_START + days_in_month
    NORMAL_COL = REMARK_COL + 1
    OVERTIME_COL = NORMAL_COL + 1
    TOTAL_COL = OVERTIME_COL + 1

    # 标题
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=TOTAL_COL)
    c = ws.cell(row=1, column=1, value=f"驾驶员月度结算表 - {year}年{month}月")
    c.font = title_font
    c.alignment = Alignment(horizontal="center")

    # 第二行：支持服务工时记录 跨越日期列
    ws.merge_cells(start_row=2, start_column=DATE_START, end_row=2, end_column=DATE_START + days_in_month - 1)
    c = ws.cell(row=2, column=DATE_START, value="支持服务工时记录")
    c.font = title_font
    c.alignment = Alignment(horizontal="center")

    # 表头
    hr = 3
    header_labels = [
        ("序号\nNO.", 6), ("合同号", 20), ("服务人员姓名", 12),
        ("服务类别", 20), ("服务等级", 8), ("属性", 6),
    ]
    for ci, (label, width) in enumerate(header_labels, 1):
        c = ws.cell(row=hr, column=ci, value=label)
        c.font = header_font
        c.border = thin_border
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(ci)].width = width

    for d in range(1, days_in_month + 1):
        c = ws.cell(row=hr, column=DATE_START + d - 1, value=str(d))
        c.font = header_font
        c.border = thin_border
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(DATE_START + d - 1)].width = 4.5

    for label, col in [("备注", REMARK_COL), ("合计正常服务天数", NORMAL_COL),
                       ("合计额外服务天数", OVERTIME_COL), ("合计总服务\n天数", TOTAL_COL)]:
        c = ws.cell(row=hr, column=col, value=label)
        c.font = header_font
        c.border = thin_border
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = 10 if label == "备注" else 14

    # 数据行
    row = 4
    idx = 0
    for rec in records_list:
        idx += 1
        person_name = rec["person_name"]
        tstatus = rec["travel_status"]
        row_sum = round(sum(v for v in rec["days"].values()), 3)
        normal_days = row_sum if rec["attribute"] == "正常" else 0
        overtime_days = row_sum if rec["attribute"] == "加班" else 0
        total_days = row_sum

        values = [
            idx, rec["contract_no"], person_name,
            rec["service_category"], rec["service_level"], rec["attribute"],
        ]
        for ci, v in enumerate(values, 1):
            c = ws.cell(row=row, column=ci, value=v)
            c.font = cell_font
            c.border = thin_border
            c.alignment = Alignment(horizontal="center", wrap_text=True)

        # 日期数据
        for d in range(1, days_in_month + 1):
            hours = rec["days"].get(str(d), 0)
            col = DATE_START + d - 1
            if hours > 0:
                c = ws.cell(row=row, column=col, value=hours)
                c.font = cell_font
            c = ws.cell(row=row, column=col)
            c.border = thin_border
            c.alignment = Alignment(horizontal="center")

        # 备注
        c = ws.cell(row=row, column=REMARK_COL, value=rec["remark"])
        c.font = cell_font
        c.border = thin_border
        c.alignment = Alignment(horizontal="center")

        # 合计列
        for col, val in [(NORMAL_COL, round(normal_days, 3)), (OVERTIME_COL, round(overtime_days, 3)), (TOTAL_COL, round(total_days, 3))]:
            c = ws.cell(row=row, column=col, value=val)
            c.font = Font(name="Arial", size=9, bold=True)
            c.border = thin_border
            c.alignment = Alignment(horizontal="center")

        row += 1

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"驾驶员月度结算-{year_month}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


# ==================== 查询接口 ====================

@router.get("/query", summary="驾驶员月度结算查询")
async def get_driver_settlement(
    year_month: str = Query(..., description="结算月份(YYYY-MM)"),
    test_order_id: int = Query(None, description="试验单号ID"),
):
    year, month = int(year_month[:4]), int(year_month[5:7])
    days_in_month = calendar.monthrange(year, month)[1]

    month_start = date(year, month, 1)
    month_end = date(year + 1, 1, 1) if month == 12 else date(year, month + 1, 1)

    from tortoise.expressions import Q
    q = (
        Q(record_date__gte=str(month_start), record_date__lt=str(month_end))
        & Q(person_type="驾驶员")
        & ~Q(approval_status="驳回")
    )
    if test_order_id and isinstance(test_order_id, int) and test_order_id > 0:
        q &= Q(test_order_id=test_order_id)

    records = await DailyRecord.filter(q).select_related("test_order").order_by("person_name", "record_date")
    records_list, person_summary = await _build_grouped_data(records, days_in_month)

    summary_out = {}
    for key, s in person_summary.items():
        # key is tuple (person_name, travel_status)
        normal_days = round(s["normal_hours"] / 8, 3) if s["normal_hours"] else 0
        overtime_days = round(s["overtime_hours"] / 8, 3) if s["overtime_hours"] else 0
        summary_key = f"{key[0]}|{key[1]}"
        summary_out[summary_key] = {
            "normal_days": normal_days,
            "overtime_days": overtime_days,
            "total_days": round(normal_days + overtime_days, 3),
            "contract_no": s["contract_no"],
        }

    return Success(data={
        "days_in_month": days_in_month,
        "records": records_list,
        "person_summary": summary_out,
    })
