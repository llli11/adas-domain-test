"""工程师月度结算"""
from decimal import Decimal
from datetime import date, datetime
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Query, Body
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.controllers.expense import monthly_settlement_controller
from app.models.expense import DailyRecord, MonthlySettlement, TestOrder
from app.schemas.base import Success
from app.schemas.expense import MonthlySettlementCreate, MonthlySettlementUpdate

router = APIRouter(tags=["费用管理"])

SUPPLIER_RATES = {
    "达安": {"local": 130, "trip": 130},
    "育喆": {"local": 36.25, "trip": 44.50},
    "万嘉禾": {"local": 35.13, "trip": 44.50},
}


@router.get("/query", summary="工程师月度考勤汇总")
async def query_engineer_attendance(year_month: str = Query(...)):
    import calendar
    y, m = int(year_month[:4]), int(year_month[5:7])
    days_in_month = calendar.monthrange(y, m)[1]
    month_start = date(y, m, 1)
    month_end = date(y + 1, 1, 1) if m == 12 else date(y, m + 1, 1)
    from tortoise.expressions import Q
    records = await DailyRecord.filter(
        Q(record_date__gte=str(month_start)) & Q(record_date__lt=str(month_end))
        & Q(person_type="工程师") & ~Q(approval_status="驳回")
    ).prefetch_related("project", "test_order").order_by("person_name", "record_date")
    groups = {}
    for rec in records:
        key = (rec.person_name, rec.test_order_id or 0, rec.project_id or 0)
        if key not in groups:
            groups[key] = {"person_name": rec.person_name, "daily": {}, "total_h": 0.0, "series": "", "tno": "", "project_name": "", "responsible": ""}
            if rec.project:
                groups[key]["series"] = rec.project.series_name or ""
                groups[key]["project_name"] = rec.project.project_name or ""
            if rec.test_order:
                groups[key]["tno"] = rec.test_order.test_order_no or ""
                groups[key]["responsible"] = rec.test_order.responsible_person or ""
        d = rec.record_date.day
        h = float(rec.work_hours or 0)
        groups[key]["daily"][str(d)] = round(groups[key]["daily"].get(str(d), 0) + h, 1)
        groups[key]["total_h"] = round(groups[key]["total_h"] + h, 1)
    data = []
    for g in groups.values():
        data.append({"person_name": g["person_name"], "series": g["series"], "test_order_no": g["tno"],
                      "project_name": g["project_name"], "responsible": g["responsible"],
                      "days": g["daily"], "total_hours": g["total_h"], "_index": len(data) + 1})
    return Success(data={"days_in_month": days_in_month, "records": data})


@router.post("/generate", summary="生成工程师月度结算")
async def generate_monthly_settlement(year_month: str = Body(..., description="结算月份(YYYY-MM)")):
    """根据该月的每日记录自动生成月度结算"""
    records = await DailyRecord.filter(
        record_date__year=int(year_month[:4]),
        record_date__month=int(year_month[5:7]),
    ).exclude(approval_status="驳回")
    test_order_groups = {}
    for r in records:
        key = r.test_order_id or 0
        if key not in test_order_groups:
            test_order_groups[key] = {"labor_cost": Decimal("0"), "advance_payment": Decimal("0")}
        test_order_groups[key]["labor_cost"] += r.work_hours * Decimal("50")
        test_order_groups[key]["advance_payment"] += r.advance_payment

    for test_order_id, costs in test_order_groups.items():
        if test_order_id == 0:
            continue
        total = costs["labor_cost"] + costs["advance_payment"]
        exist = await MonthlySettlement.filter(
            test_order_id=test_order_id, year_month=year_month
        ).first()
        if not exist:
            await monthly_settlement_controller.create(obj_in=MonthlySettlementCreate(
                test_order_id=test_order_id,
                year_month=year_month,
                labor_cost=float(costs["labor_cost"]),
                advance_payment=float(costs["advance_payment"]),
                total_amount=float(total),
            ))
    return Success(msg=f"已生成 {year_month} 月度结算")


@router.get("/list", summary="查看工程师月度结算列表")
async def list_monthly_settlements(
    year_month: str = Query(None, description="结算月份"),
    status: str = Query(None, description="状态"),
    test_order_id: int = Query(None, description="试验单号ID"),
):
    items = await monthly_settlement_controller.search(
        year_month=year_month, status=status, test_order_id=test_order_id
    )
    data = []
    for item in items:
        d = await item.to_dict()
        d["test_order_no"] = item.test_order.test_order_no if item.test_order else ""
        data.append(d)
    return Success(data=data)


@router.post("/update", summary="更新工程师月度结算")
async def update_monthly_settlement(
    id: int = Query(...),
    item_in: MonthlySettlementUpdate = None,
):
    await monthly_settlement_controller.update(id=id, obj_in=item_in)
    return Success(msg="更新成功")


@router.get("/export", summary="导出工程师月度结算单Excel")
async def export_monthly_settlement(year_month: str = Query(...)):
    """导出界面显示的数据到Excel"""
    import calendar
    y, m = int(year_month[:4]), int(year_month[5:7])
    days_in_month = calendar.monthrange(y, m)[1]
    month_start = date(y, m, 1)
    month_end = date(y + 1, 1, 1) if m == 12 else date(y, m + 1, 1)
    from tortoise.expressions import Q
    records = await DailyRecord.filter(
        Q(record_date__gte=str(month_start)) & Q(record_date__lt=str(month_end))
        & Q(person_type="工程师") & ~Q(approval_status="驳回")
    ).prefetch_related("project", "test_order").order_by("person_name", "record_date")

    groups = {}
    for rec in records:
        key = (rec.person_name, rec.test_order_id or 0, rec.project_id or 0)
        if key not in groups:
            groups[key] = {"person_name": rec.person_name, "daily": {}, "total_h": 0.0, "project_name": "", "responsible": "", "tno": ""}
            if rec.project: groups[key]["project_name"] = rec.project.project_name or ""
            if rec.test_order:
                groups[key]["responsible"] = rec.test_order.responsible_person or ""
                groups[key]["tno"] = rec.test_order.test_order_no or ""
        d = rec.record_date.day
        h = float(rec.work_hours or 0)
        groups[key]["daily"][str(d)] = round(groups[key]["daily"].get(str(d), 0) + h, 1)
        groups[key]["total_h"] = round(groups[key]["total_h"] + h, 1)

    wb = Workbook()
    ws = wb.active
    ws.title = f"{y}年{m}月工程师结算"

    title_font = Font(name="Arial", size=12, bold=True)
    header_font = Font(name="Arial", size=9, bold=True)
    cell_font = Font(name="Arial", size=9)
    thin = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

    # Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=days_in_month + 6)
    ws.cell(row=1, column=1, value=f"工程师月度结算 - {y}年{m}月").font = title_font
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    # Headers
    hr = 2
    headers = ["序号", "姓名", "专业", "属性"] + [str(d) for d in range(1, days_in_month + 1)] + ["合计总工时", "备注"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=hr, column=ci, value=h)
        c.font = header_font; c.border = thin; c.alignment = Alignment(horizontal="center", wrap_text=True)

    row = 3
    for i, g in enumerate(groups.values()):
        ws.cell(row=row, column=1, value=i + 1).font = cell_font
        ws.cell(row=row, column=2, value=g["person_name"]).font = cell_font
        ws.cell(row=row, column=3, value="智能驾驶专项").font = cell_font
        prop = f'{g["project_name"]}\n{g["responsible"]}'
        ws.cell(row=row, column=4, value=prop).font = cell_font
        for d in range(1, days_in_month + 1):
            v = g["daily"].get(str(d), 0)
            if v > 0:
                ws.cell(row=row, column=4 + d, value=v).font = cell_font
        ws.cell(row=row, column=4 + days_in_month + 1, value=g["total_h"]).font = Font(name="Arial", size=9, bold=True)
        ws.cell(row=row, column=4 + days_in_month + 2, value=g.get("tno", "")).font = cell_font
        for ci in range(1, days_in_month + 7):
            ws.cell(row=row, column=ci).border = thin
        row += 1

    # Column widths
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 24
    for d in range(1, days_in_month + 1):
        ws.column_dimensions[get_column_letter(4 + d)].width = 5
    ws.column_dimensions[get_column_letter(4 + days_in_month + 1)].width = 12
    ws.column_dimensions[get_column_letter(4 + days_in_month + 2)].width = 16

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"工程师月度结算-{year_month}.xlsx"
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})
